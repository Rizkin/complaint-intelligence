"""Build the 200-complaint label-noise audit sheet (xlsx with dropdowns) from the held-out sample.

The sheet hides model predictions on purpose. Output goes to data/eval/ (gitignored: it holds narratives).
After labeling, `evals/classify.py report-audit` (to be added) reads the filled sheet.

Usage: uv run python evals/make_audit_sheet.py [--out data/eval/audit_200.xlsx]
"""

import argparse
import json

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ap = argparse.ArgumentParser()
ap.add_argument("--out", default="data/eval/audit_200.xlsx")
ap.add_argument("--per-product", type=int, default=100)
args = ap.parse_args()

df = pd.read_parquet("data/eval/heldout_2026_05.parquet")
labels = json.load(open("data/eval/labels.json"))
parts = [g.sample(n=args.per_product, random_state=7) for _, g in df.groupby("product")]
audit = pd.concat(parts).sample(frac=1, random_state=7).reset_index(drop=True)  # shuffled so products are mixed

wb = Workbook()
ws = wb.active
ws.title = "Audit"
cols = ["#", "complaint_id", "narrative", "consumer_product", "consumer_issue",
        "product_correct (Y/N)", "issue_correct (Y/N/Unsure)", "better_issue (if N)", "note"]
ws.append(cols)
for i, r in enumerate(audit.itertuples(), 1):
    ws.append([i, r.complaint_id, r.narrative, r.product, r.issue, None, None, None, None])

head = PatternFill("solid", fgColor="1F4E78")
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = head
    c.alignment = Alignment(wrap_text=True, vertical="center")
widths = {"A": 5, "B": 12, "C": 90, "D": 24, "E": 34, "F": 14, "G": 16, "H": 36, "I": 30}
for k, v in widths.items():
    ws.column_dimensions[k].width = v
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "D2"

lab = wb.create_sheet("Labels")
lab.append(["all_issues"])
for x in labels["all_issues"]:
    lab.append([x])
n_issue = len(labels["all_issues"]) + 1

dv_yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
dv_yns = DataValidation(type="list", formula1='"Y,N,Unsure"', allow_blank=True)
dv_issue = DataValidation(type="list", formula1=f"=Labels!$A$2:$A${n_issue}", allow_blank=True)
for dv in (dv_yn, dv_yns, dv_issue):
    ws.add_data_validation(dv)
last = len(audit) + 1
dv_yn.add(f"F2:F{last}")
dv_yns.add(f"G2:G{last}")
dv_issue.add(f"H2:H{last}")

guide = wb.create_sheet("Guide", 0)
for line in [
    "Label-noise audit (plan: Thu Oct 15, 3 h). 200 complaints, held-out month 2026-05, 100 per product.",
    "For each row read the narrative and judge the CONSUMER's labels, not a model's:",
    "  product_correct: does the complaint really concern a checking/savings account (vs a credit card)? Y or N.",
    "  issue_correct: is the consumer-chosen issue the best fit among the CFPB issues? Y / N / Unsure.",
    "  better_issue: if N, pick the issue you think fits best (dropdown, all 24 labels on the Labels sheet).",
    "  note: optional, e.g. 'mentions both', 'unclear', 'spanish'.",
    "Rules: judge on the narrative only. Text like XXXX is CFPB redaction. If the narrative is too vague, mark Unsure.",
    "Save the file in place. Target pace: ~50 seconds per row.",
    "Model predictions are hidden on purpose so they cannot bias you.",
]:
    guide.append([line])
guide.column_dimensions["A"].width = 130

wb.save(args.out)
print(f"wrote {args.out}: {len(audit)} rows ({audit.groupby('product').size().to_dict()})")

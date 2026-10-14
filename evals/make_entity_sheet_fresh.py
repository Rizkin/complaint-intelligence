"""Build a FRESH entity-matching test sheet: companies not in the earlier 100-pair sheet and not in the overrides seed.

Predictions come from the final matcher table (data/eval/entity_match.csv). Same columns/dropdowns as the 100-pair sheet so
evals/eval_entity_match.py works unchanged (use --sheet data/eval/entity_pairs_fresh50.xlsx).
Composition (seed 23): 15 high-volume, 15 low-volume matched, 10 rejected-but-bank-like, 10 other rejected.
Usage: uv run python evals/make_entity_sheet_fresh.py
"""

import json
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

df = pd.read_csv("data/eval/entity_match.csv")
seen = set(pd.read_excel("data/eval/entity_pairs_100.xlsx", sheet_name="Pairs").cfpb_company)
seen |= set(pd.read_csv("pipelines/dbt/seeds/entity_overrides.csv").cfpb_company)
seen |= set(pd.read_csv("pipelines/dbt/seeds/entity_map_seed.csv").cfpb_company)
df = df[~df.company.isin(seen)].copy()
df["alts"] = df.alt_candidates.map(json.loads)
bank_like = re.compile(r"BANK|BANCORP|BANCSHARES|TRUST|SAVINGS|FEDERAL|FINANCIAL|BANC|CAPITAL", re.I)
big = df[df.n_complaints >= 100]
tail_m = df[(df.n_complaints < 100) & df.cert.notna()]
rej_bl = df[df.cert.isna() & df.company.map(lambda x: bool(bank_like.search(x)))]
rej_other = df[df.cert.isna() & ~df.company.map(lambda x: bool(bank_like.search(x))) & (df.n_complaints < 100)]
parts = [big.sample(n=min(15, len(big)), random_state=23), tail_m.sample(n=min(15, len(tail_m)), random_state=23),
         rej_bl.sample(n=min(10, len(rej_bl)), random_state=23), rej_other.sample(n=min(10, len(rej_other)), random_state=23)]
pick = pd.concat(parts).drop_duplicates("company")
if len(pick) < 50:
    pick = pd.concat([pick, df[~df.company.isin(pick.company)].sample(n=50 - len(pick), random_state=23)])
pick = pick.sample(frac=1, random_state=23).reset_index(drop=True)

wb = Workbook()
ws = wb.active
ws.title = "Pairs"
cols = ["#", "cfpb_company", "complaints", "predicted_cert", "predicted_bank", "state", "holding_company", "assets_bn", "score",
        "alt_1 (cert | name | score)", "alt_2", "alt_3", "verdict", "correct_cert (if 'Wrong, different bank')", "note", "broad_partner_certs (not ownership)"]
ws.append(cols)


def alt(r, i):
    al = [x for x in r.alts if x["cert"] != r.cert]
    return f'{al[i]["cert"]} | {al[i]["name"]} | {al[i]["score"]}' if len(al) > i else ""


for i, r in enumerate(pick.itertuples(), 1):
    ws.append([i, r.company, int(r.n_complaints), None if pd.isna(r.cert) else int(r.cert),
               r.institution if pd.notna(r.institution) else "(no bank predicted)", r.state if pd.notna(r.state) else "",
               r.holding_company if pd.notna(r.holding_company) else "", None if pd.isna(r.assets_bn) else float(r.assets_bn),
               float(r.score), alt(r, 0), alt(r, 1), alt(r, 2), None, None, None, None])
fill = PatternFill("solid", fgColor="1F4E78")
for c in ws[1]:
    c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), fill, Alignment(wrap_text=True, vertical="center")
dv = DataValidation(type="list", formula1='"Right,Wrong - different bank,Wrong - no FDIC bank exists,Unsure"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"M2:M{len(pick) + 1}")
wb.save("data/eval/entity_pairs_fresh50.xlsx")
print(f"wrote data/eval/entity_pairs_fresh50.xlsx: {len(pick)} rows; predicted-match {int(pick.cert.notna().sum())}, no-bank {int(pick.cert.isna().sum())}; "
      f"{pick.n_complaints.sum():,} complaints")
